from tkinter import *
from tkinter.ttk import Combobox
import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl import Workbook
import pathlib
import os
from PIL import Image, ImageTk

root = Tk()
root.title("Data Entry")
root.geometry('702x400+300+200')
root.resizable(False, False)
root.configure(bg="#326273")

# Get script base directory for reliable pathing
script_dir = os.path.dirname(os.path.abspath(__file__))
excel_file_path = os.path.join(script_dir, 'Backened_data.xlsx')

# Check and create Excel file if it doesn't exist
file_path_obj = pathlib.Path(excel_file_path)
if not file_path_obj.exists():
    workbook = Workbook()
    sheet = workbook.active
    sheet['A1'] = "Full Name"
    sheet['B1'] = "phoneNumber"
    sheet['C1'] = "Age"
    sheet['D1'] = "Gender"
    sheet['E1'] = "Address"
    workbook.save(excel_file_path)

def submit():
    name = namevalue.get()
    Contact = Contactvalue.get()
    age = Agevalue.get()
    gender = gender_combobox.get()
    address = addressEntry.get(1.0, END).strip()

    if not name or not Contact:
        messagebox.showwarning("Warning", "Name and Contact Number are required!")
        return

    try:
        file = openpyxl.load_workbook(excel_file_path)
        sheet = file.active
        
        row_num = sheet.max_row + 1
        sheet.cell(column=1, row=row_num, value=name)
        sheet.cell(column=2, row=row_num, value=Contact)
        sheet.cell(column=3, row=row_num, value=age)
        sheet.cell(column=4, row=row_num, value=gender)
        sheet.cell(column=5, row=row_num, value=address)

        file.save(excel_file_path)
        messagebox.showinfo('info', 'Detail added!')

        clear()
    except Exception as e:
        messagebox.showerror("Error", f"Could not save data: {e}")

def clear():
    namevalue.set('')
    Contactvalue.set('')
    Agevalue.set('')
    gender_combobox.set('Male')
    addressEntry.delete(1.0, END)

# Icon Configuration using Pillow for reliable PNG loading and rendering
try:
    icon_path = os.path.join(script_dir, "logo.png")
    img = Image.open(icon_path)
    root.icon_image = ImageTk.PhotoImage(img)
    root.iconphoto(False, root.icon_image)
except Exception as e:
    print("Icon load error:", e)

# Heading
Label(root, text="Please fill out this Entry form:", font="arial 13", bg="#326273", fg="#fff").place(x=20, y=20)

# Labels
Label(root, text='Name', font=23, bg="#326273", fg="#fff").place(x=50, y=100)
Label(root, text='Contact No', font=23, bg="#326273", fg="#fff").place(x=50, y=150)
Label(root, text='Age', font=23, bg="#326273", fg="#fff").place(x=50, y=200)
Label(root, text='Gender', font=23, bg="#326273", fg="#fff").place(x=370, y=200)
Label(root, text='Address', font=23, bg="#326273", fg="#fff").place(x=50, y=250)

# Entry Variables
namevalue = StringVar()
Contactvalue = StringVar()
Agevalue = StringVar()

nameEntry = Entry(root, textvariable=namevalue, width=45, bd=2, font=20)
contactEntry = Entry(root, textvariable=Contactvalue, width=45, bd=2, font=20)
ageEntry = Entry(root, textvariable=Agevalue, width=15, bd=2, font=20)

# Gender
gender_combobox = Combobox(root, values=['Male', 'Female', 'Other'], font='arial 14', state='readonly', width=14)
gender_combobox.place(x=440, y=200)
gender_combobox.set('Male')

addressEntry = Text(root, width=50, height=4, bd=4)

nameEntry.place(x=200, y=100)
contactEntry.place(x=200, y=150)
ageEntry.place(x=200, y=200)
addressEntry.place(x=200, y=250)

Button(root, text="Submit", bg="#326273", fg="white", width=15, height=2, command=submit).place(x=200, y=350)
Button(root, text="Clear", bg="#326273", fg="white", width=15, height=2, command=clear).place(x=340, y=350)
Button(root, text="Exit", bg="#326273", fg="white", width=15, height=2, command=root.destroy).place(x=480, y=350)

root.mainloop()